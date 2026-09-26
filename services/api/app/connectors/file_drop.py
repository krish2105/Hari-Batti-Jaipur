"""CSV / Excel exports dropped into a watched folder (an SFTP sync job can fill the folder).

Each new or changed file is read once; every row is one vendor record. Files are never modified,
moved or deleted: the folder belongs to the police/vendor, we only read it.
"""

import asyncio
import csv
from collections.abc import AsyncIterator
from pathlib import Path
from typing import Any

from .base import Mapping, ReadOnlyConnector

SUFFIXES = (".csv", ".xlsx")


def read_table(path: Path) -> list[dict[str, Any]]:
    """Rows of a CSV or the first sheet of an .xlsx file, as dicts keyed by the header row."""
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            return list(csv.DictReader(f))
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    head = [str(h).strip() if h is not None else "" for h in rows[0]]
    return [dict(zip(head, r, strict=False)) for r in rows[1:] if any(v is not None for v in r)]


class FileDropConnector(ReadOnlyConnector):
    kind = "file_drop"

    def __init__(self, connector_id: str, mapping: Mapping, folder: str | Path, every_s: float = 5.0):
        super().__init__(connector_id, mapping)
        self.folder, self.every_s = Path(folder), every_s
        self._seen: dict[Path, float] = {}

    def _new_files(self) -> list[Path]:
        """Files that are new or changed since the last look (oldest first)."""
        if not self.folder.is_dir():
            return []
        out = []
        for p in sorted(self.folder.iterdir(), key=lambda p: p.stat().st_mtime):
            if p.suffix.lower() in SUFFIXES and self._seen.get(p) != p.stat().st_mtime:
                self._seen[p] = p.stat().st_mtime
                out.append(p)
        return out

    async def _read(self) -> AsyncIterator[Any]:
        while True:
            for p in self._new_files():
                try:
                    yield await asyncio.to_thread(read_table, p)
                except (OSError, ValueError, KeyError):
                    self.rejected += 1
            await asyncio.sleep(self.every_s)
